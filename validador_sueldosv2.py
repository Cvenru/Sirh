import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

# CONFIGURACIÓN: Modificar según tu archivo
archivo = 'ARCHIVO_SIRH.xlsx'  # Cambiar por tu archivo
hoja = 'Hoja3'  # Cambiar por tu hoja

df = pd.read_excel(archivo, sheet_name=hoja)
print(df)

print("Columnas disponibles:")
print(df.columns.tolist())

########################### Validación de tipos de pagos
# Lista de valores válidos (modificar según necesidad)
valores_validos = [
    'PAGO HONORARIO',
    'PAGO NORMAL',
    'PAGO ACCESORIO-BONO VACACIONES',
    'PAGO ACCESORIO-BONO MENSUAL',
    'PAGO ACCESORIO RETROACTIVO',
    'PAGO ACCESORIO RETROACTIVO-DIF.NORMAL.REL.L.M.',
    'PAGO ACCESORIO RETROACTIVO-PAGO ACCESORIO',
    'PAGO ACCESORIO RETROACTIVO-DIF.BON.ENFERMERAS',
    'PAGO ACCESORIO-PAGO ACCESORIO',
    'PAGO ACCESORIO RETROACTIVO-DIF. POR ASCENSOS'
]

unicos = df['PROCESO'].unique()
valores_raros = [v for v in unicos if v not in valores_validos]

print(f"Únicos: {len(unicos)} | Válidos: {len(valores_validos)} | Raros: {len(valores_raros)}")

if valores_raros:
    print(f"\n⚠️ Valores raros encontrados:")
    print(valores_raros)
else:
    print("\n✓ Todos los procesos son válidos")


########################### Validación de horas extras
horas_extra = 163  # MODIFICAR según criterio

df_horas_extras = df[df['CANT. HRS. EXTRAS'] > horas_extra]
if len(df_horas_extras) > 0:
    print(f"\n⚠️ ALERTA: {len(df_horas_extras)} registros con HORAS > {horas_extra}")
    print("\nDetalle:")
    print(df_horas_extras[['Nombre', 'Identificación', 'CANT. HRS. EXTRAS', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS','CONTRATO CORTO']])
else:
    print(f"\n✓ No hay registros con HORAS > {horas_extra}")


############################ Definición de plantas
planta_administrativo = ['ADMINISTRATIVOS', 'AUXILIARES', 'TECNICOS']
planta_profesional = ['PROFESIONALES', 'BIOQUIMICOS', 'QUIMICOS', 'QUÍMICOS FARMACÉUTICOS']
planta_medica = ['MEDICOS', 'ODONTOLOGOS']
planta_directiva = ['DIRECTIVOS']

############################ Validación planta administrativa
sueldo_max_admin = 0 #Modificar según criterio

df_admin_exceso = df[
    (df['PLANTA '].isin(planta_administrativo)) & 
    (df['PROCESO'].isin(['PAGO HONORARIO', 'PAGO NORMAL'])) &
    (df['Salario Base'] > sueldo_max_admin)
]

if len(df_admin_exceso) > 0:
    print(f"\n⚠️ ALERTA: {len(df_admin_exceso)} registros ADMINISTRATIVOS con SALARIO BASE > {sueldo_max_admin:,}")
    print("\nDetalle:")
    df_admin_exceso['DIFERENCIA'] = df_admin_exceso['Salario Base'] - sueldo_max_admin
    print(df_admin_exceso[['PROCESO', 'Nombre', 'Identificación', 'PLANTA ', 'Salario Base', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS', 'CONTRATO CORTO', 'DIFERENCIA']])
else:
    print(f"\n✓ No hay registros ADMINISTRATIVOS con SALARIO BASE > {sueldo_max_admin:,}")


############################ Validación planta profesional
sueldo_max_profesional = 0 # Modificar según criterio

df_profesional_exceso = df[
    (df['PLANTA '].isin(planta_profesional)) & 
    (df['PROCESO'].isin(['PAGO HONORARIO', 'PAGO NORMAL'])) &
    (df['Salario Base'] > sueldo_max_profesional)
]

if len(df_profesional_exceso) > 0:
    print(f"\n⚠️ ALERTA: {len(df_profesional_exceso)} registros PROFESIONALES con SALARIO BASE > {sueldo_max_profesional:,}")
    print("\nDetalle:")
    df_profesional_exceso['DIFERENCIA'] = df_profesional_exceso['Salario Base'] - sueldo_max_profesional
    print(df_profesional_exceso[['PROCESO', 'Nombre', 'Identificación', 'PLANTA ', 'Salario Base', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS', 'CONTRATO CORTO', 'DIFERENCIA']])
else:
    print(f"\n✓ No hay registros PROFESIONALES con SALARIO BASE > {sueldo_max_profesional:,}")

############################ Validación planta médica
sueldo_max_medica = 0 # Modificar según criterio
df_medica_exceso = df[
    (df['PLANTA '].isin(planta_medica)) & 
    (df['PROCESO'].isin(['PAGO HONORARIO', 'PAGO NORMAL'])) &
    (df['Salario Base'] > sueldo_max_medica)
]

if len(df_medica_exceso) > 0:
    print(f"\n⚠️ ALERTA: {len(df_medica_exceso)} registros MEDICA con SALARIO BASE > {sueldo_max_medica:,}")
    print("\nDetalle:")
    df_medica_exceso['DIFERENCIA'] = df_medica_exceso['Salario Base'] - sueldo_max_medica
    print(df_medica_exceso[['PROCESO', 'Nombre', 'Identificación', 'PLANTA ', 'Salario Base', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS', 'CONTRATO CORTO', 'DIFERENCIA']])
else:
    print(f"\n✓ No hay registros MEDICA con SALARIO BASE > {sueldo_max_medica:,}")

############################ Validación planta directiva
sueldo_min_directiva = 0 # Modificar según criterio

df_directiva_bajo = df[
    (df['PLANTA '].isin(planta_directiva)) & 
    (df['PROCESO'].isin(['PAGO HONORARIO', 'PAGO NORMAL'])) &
    (df['Salario Base'] < sueldo_min_directiva)
]

if len(df_directiva_bajo) > 0:
    print(f"\n⚠️ ALERTA: {len(df_directiva_bajo)} registros DIRECTIVA con SALARIO BASE < {sueldo_min_directiva:,}")
    print("\nDetalle:")
    df_directiva_bajo['DIFERENCIA'] = sueldo_min_directiva - df_directiva_bajo['Salario Base']
    print(df_directiva_bajo[['PROCESO', 'Nombre', 'Identificación', 'PLANTA ', 'Salario Base', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS', 'CONTRATO CORTO', 'DIFERENCIA']])
else:
    print(f"\n✓ No hay registros DIRECTIVA con SALARIO BASE < {sueldo_min_directiva:,}")


############################ Validación bonificaciones
df_bono_exceso = df[df['Beneficios Laborales'] > df['Salario Base']]

if len(df_bono_exceso) > 0:
    print(f"\n⚠️ ALERTA: {len(df_bono_exceso)} registros con BONIFICACIÓN > SALARIO BASE")
    print("\nDetalle:")
    print(df_bono_exceso[['PROCESO', 'Nombre', 'Identificación', 'Salario Base', 'Beneficios Laborales', 'UNIDAD', 'LEY', 'DIAS TRABAJADOS', 'CONTRATO CORTO']])
else:
    print(f"\n✓ No hay registros con BONIFICACIÓN > SALARIO BASE")


############################ Resumen de gastos
print("\nResumen de gastos por tipo de proceso:")
resumen_proceso = df.groupby('PROCESO')['Salario Base'].sum()
print(resumen_proceso)

total_salario_base = df['Salario Base'].sum()
total_bonificaciones = df['Beneficios Laborales'].sum()
total_horas_extras = df['MTO.HRS.EXTRAS'].sum()
print(f"\nTotal Salario Base: ${total_salario_base:,.0f}")
print(f"Total Bonificaciones: ${total_bonificaciones:,.0f}")
print(f"Total MTO. HRS. EXTRAS: ${total_horas_extras:,.0f}")
print(f"Gran Total: ${total_salario_base + total_bonificaciones + total_horas_extras:,.0f}")

############################ Clasificar plantas
def clasificar_planta(planta):
    if planta in planta_administrativo:
        return 'Administrativa'
    elif planta in planta_profesional:
        return 'Profesional'
    elif planta in planta_medica:
        return 'Médica'
    elif planta in planta_directiva:
        return 'Directiva'
    else:
        return 'Otra'

df['Tipo Planta'] = df['PLANTA '].apply(clasificar_planta)

############################ Generar reporte Excel
with pd.ExcelWriter('reporte_completo.xlsx', engine='openpyxl') as writer:
    # 1. Resumen general
    resumen = pd.DataFrame({
        'Concepto': ['Salario Base', 'Bonificaciones', 'Horas Extras', 'TOTAL'],
        'Monto': [total_salario_base, total_bonificaciones, total_horas_extras, 
                  total_salario_base + total_bonificaciones + total_horas_extras]
    })
    resumen.to_excel(writer, sheet_name='Resumen', index=False)
    
    # 2. Gastos por Proceso
    resumen_proceso_df = pd.DataFrame({
        'Proceso': resumen_proceso.index,
        'Total': resumen_proceso.values
    })
    resumen_proceso_df.to_excel(writer, sheet_name='Gastos_Proceso', index=False)
    
    # 3. Gastos por Planta
    resumen_planta = df.groupby('Tipo Planta')['Salario Base'].sum()
    resumen_planta_df = pd.DataFrame({
        'Tipo Planta': resumen_planta.index,
        'Total Salario': resumen_planta.values
    })
    resumen_planta_df.to_excel(writer, sheet_name='Gastos_Planta', index=False)
    
    # 4. Alertas
    if len(df_admin_exceso) > 0:
        df_admin_exceso.to_excel(writer, sheet_name='Admin_Exceso', index=False)
    if len(df_profesional_exceso) > 0:
        df_profesional_exceso.to_excel(writer, sheet_name='Prof_Exceso', index=False)
    if len(df_medica_exceso) > 0:
        df_medica_exceso.to_excel(writer, sheet_name='Med_Exceso', index=False)
    if len(df_directiva_bajo) > 0:
        df_directiva_bajo.to_excel(writer, sheet_name='Dir_Bajo', index=False)
    if len(df_horas_extras) > 0:
        df_horas_extras.to_excel(writer, sheet_name='Horas_Extras', index=False)
    if len(df_bono_exceso) > 0:
        df_bono_exceso.to_excel(writer, sheet_name='Bono_Exceso', index=False)

############################ Agregar gráficos
wb = load_workbook('reporte_completo.xlsx')

# Gráfico 1: Gastos por Proceso
ws1 = wb['Gastos_Proceso']
chart1 = BarChart()
chart1.title = "Gastos por Tipo de Proceso"
chart1.y_axis.title = "Total ($)"
chart1.x_axis.title = "Proceso"
data1 = Reference(ws1, min_col=2, min_row=1, max_row=len(resumen_proceso_df) + 1)
cats1 = Reference(ws1, min_col=1, min_row=2, max_row=len(resumen_proceso_df) + 1)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
ws1.add_chart(chart1, "D2")

# Gráfico 2: Gastos por Planta
ws2 = wb['Gastos_Planta']
chart2 = BarChart()
chart2.title = "Gastos por Tipo de Planta"
chart2.y_axis.title = "Total Salario ($)"
chart2.x_axis.title = "Tipo de Planta"
data2 = Reference(ws2, min_col=2, min_row=1, max_row=len(resumen_planta_df) + 1)
cats2 = Reference(ws2, min_col=1, min_row=2, max_row=len(resumen_planta_df) + 1)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws2.add_chart(chart2, "D2")

wb.save('reporte_completo.xlsx')
print("\n✅ Reporte Excel con gráficos guardado")